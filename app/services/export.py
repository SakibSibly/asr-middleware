"""
Export service for transcript formats
"""

import json
from typing import List
from datetime import datetime
import os
import tempfile

from app.models import Meeting, TranscriptSegment


class ExportService:
    """Handle transcript exports in various formats"""
    
    def export_txt(
        self,
        segments: List[TranscriptSegment],
        include_translations: bool = True,
        include_timestamps: bool = True,
    ) -> str:
        """Export as plain text"""
        lines = []
        
        for segment in segments:
            line_parts = []
            
            # Add timestamp
            if include_timestamps:
                timestamp = self._format_timestamp(segment.start_time)
                line_parts.append(f"[{timestamp}]")
            
            # Add speaker
            if segment.speaker_name:
                line_parts.append(f"{segment.speaker_name}:")
            
            # Add text
            if include_translations and segment.translated_text:
                line_parts.append(segment.translated_text)
            else:
                line_parts.append(segment.original_text)
            
            lines.append(" ".join(line_parts))
        
        return "\n".join(lines)
    
    def export_json(
        self,
        meeting: Meeting,
        segments: List[TranscriptSegment],
        include_translations: bool = True,
    ) -> str:
        """Export as JSON"""
        data = {
            "meeting": {
                "id": meeting.id,
                "title": meeting.title,
                "date_time": meeting.date_time.isoformat(),
                "duration": meeting.duration,
                "participants": meeting.participants,
            },
            "segments": [],
        }
        
        for segment in segments:
            segment_data = {
                "id": segment.id,
                "speaker": segment.speaker_name,
                "start_time": segment.start_time,
                "end_time": segment.end_time,
                "original_text": segment.original_text,
                "original_language": segment.original_language,
                "confidence_score": segment.confidence_score,
            }
            
            if include_translations and segment.translated_text:
                segment_data["translated_text"] = segment.translated_text
                segment_data["translation_service"] = segment.translation_service
            
            data["segments"].append(segment_data)
        
        return json.dumps(data, indent=2)
    
    def export_srt(
        self,
        segments: List[TranscriptSegment],
        include_translations: bool = True,
    ) -> str:
        """Export as SRT subtitles"""
        lines = []
        
        for i, segment in enumerate(segments, 1):
            # Subtitle number
            lines.append(str(i))
            
            # Timestamp
            start = self._format_srt_timestamp(segment.start_time)
            end = self._format_srt_timestamp(segment.end_time)
            lines.append(f"{start} --> {end}")
            
            # Text
            text = segment.translated_text if include_translations and segment.translated_text else segment.original_text
            lines.append(text)
            
            # Blank line
            lines.append("")
        
        return "\n".join(lines)
    
    def export_vtt(
        self,
        segments: List[TranscriptSegment],
        include_translations: bool = True,
    ) -> str:
        """Export as WebVTT subtitles"""
        lines = ["WEBVTT", ""]
        
        for segment in segments:
            # Timestamp
            start = self._format_vtt_timestamp(segment.start_time)
            end = self._format_vtt_timestamp(segment.end_time)
            lines.append(f"{start} --> {end}")
            
            # Text
            text = segment.translated_text if include_translations and segment.translated_text else segment.original_text
            if segment.speaker_name:
                lines.append(f"<v {segment.speaker_name}>{text}")
            else:
                lines.append(text)
            
            # Blank line
            lines.append("")
        
        return "\n".join(lines)
    
    def export_docx(
        self,
        meeting: Meeting,
        segments: List[TranscriptSegment],
        include_translations: bool = True,
        include_timestamps: bool = True,
    ) -> str:
        """Export as Word document"""
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            
            doc = Document()
            
            # Add title
            title = doc.add_heading(meeting.title, 0)
            
            # Add metadata
            doc.add_paragraph(f"Date: {meeting.date_time.strftime('%Y-%m-%d %H:%M')}")
            if meeting.duration:
                doc.add_paragraph(f"Duration: {meeting.duration // 60} minutes")
            if meeting.participants:
                doc.add_paragraph(f"Participants: {', '.join(meeting.participants)}")
            
            doc.add_paragraph()  # Blank line
            
            # Add transcript
            doc.add_heading("Transcript", 1)
            
            for segment in segments:
                p = doc.add_paragraph()
                
                # Add timestamp
                if include_timestamps:
                    timestamp = self._format_timestamp(segment.start_time)
                    run = p.add_run(f"[{timestamp}] ")
                    run.font.color.rgb = RGBColor(128, 128, 128)
                    run.font.size = Pt(9)
                
                # Add speaker
                if segment.speaker_name:
                    run = p.add_run(f"{segment.speaker_name}: ")
                    run.bold = True
                
                # Add text
                text = segment.translated_text if include_translations and segment.translated_text else segment.original_text
                p.add_run(text)
            
            # Save to temporary file
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
            doc.save(temp_file.name)
            
            return temp_file.name
        
        except Exception as e:
            raise Exception(f"Failed to create DOCX: {str(e)}")
    
    def _format_timestamp(self, seconds: float) -> str:
        """Format seconds as HH:MM:SS"""
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    
    def _format_srt_timestamp(self, seconds: float) -> str:
        """Format seconds as SRT timestamp (HH:MM:SS,mmm)"""
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"
    
    def _format_vtt_timestamp(self, seconds: float) -> str:
        """Format seconds as VTT timestamp (HH:MM:SS.mmm)"""
        hours = int(seconds // 3600)
        minutes = int((seconds % 3600) // 60)
        secs = int(seconds % 60)
        millis = int((seconds % 1) * 1000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"
